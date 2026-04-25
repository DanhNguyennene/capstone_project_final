# make_benchmark_ppt.py
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

HEADERS = ["Paper", "Why it matches your system", "Compare on"]

ROWS = [
    ("NL2Bash (2018)", "Closest to NL->shell command translation", "command/tool selection accuracy, exact-match/action-match"),
    ("InterCode (2023)", "Interactive Bash/SQL/Python with execution feedback", "multi-turn success, state-correct completion"),
    ("AgentBench (ICLR 2024)", "General LLM-agent benchmark in interactive environments", "routing reliability, multi-step decision quality"),
    ("ToolBench paper (2023)", "Tool-manipulation benchmark for real APIs", "tool recall, wrong-tool rate, tool-chain success"),
    ("MCP-AgentBench (2025)", "MCP-native tool benchmark (very close to your architecture)", "MCP tool success, interoperability, outcome-based success"),
    ("Agent-SafetyBench (2024/2025)", "Safety failures in tool-using agents", "HITL trigger quality, unsafe-action prevention"),
    ("R-Judge (EMNLP Findings 2024)", "Risk-awareness benchmark from interaction traces", "risk detection before tool execution"),
    ("GAP benchmark (2026)", "Text refusal vs actual harmful tool-call divergence", "\"said no but still executed\" failure rate"),
    ("Cluster Workload Allocation (IEEE Access 2026)", "Domain-near baseline for NL-driven cluster scheduling", "intent parsing quality, scheduling outcome quality"),
]

ROWS_PER_SLIDE = 5


def write_cell(cell, text, bold=False, font_size=11, color=(0, 0, 0)):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.bold = bold
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color)


def add_table_slide(prs, chunk, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = f"Benchmark Comparison ({idx}/{total})"

    x, y, w, h = Inches(0.3), Inches(1.2), Inches(12.7), Inches(5.8)
    table = slide.shapes.add_table(len(chunk) + 1, 3, x, y, w, h).table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(5.5)
    table.columns[2].width = Inches(4.7)

    for c, header in enumerate(HEADERS):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(31, 78, 121)
        write_cell(cell, header, bold=True, font_size=12, color=(255, 255, 255))

    for r, row in enumerate(chunk, start=1):
        for c, value in enumerate(row):
            write_cell(table.cell(r, c), value, font_size=11)


def main():
    prs = Presentation()
    chunks = [ROWS[i:i + ROWS_PER_SLIDE] for i in range(0, len(ROWS), ROWS_PER_SLIDE)]
    for i, chunk in enumerate(chunks, start=1):
        add_table_slide(prs, chunk, i, len(chunks))

    out_file = "benchmark_comparison.pptx"
    prs.save(out_file)
    print(f"Saved {out_file}")


if __name__ == "__main__":
    main()